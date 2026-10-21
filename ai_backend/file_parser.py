import io
import os
import base64
import logging
import pdfplumber
from docx import Document

logger = logging.getLogger("file_parser")

# 全局预加载 MarkItDown 单例，避免每次解析重复初始化
_markitdown_instance = None
try:
    from markitdown import MarkItDown
    _markitdown_instance = MarkItDown()
    print("[file_parser] [OK] Microsoft MarkItDown engine loaded and active")
except Exception as e:
    print(f"[file_parser] [WARN] MarkItDown init skipped: {e}")

def parse_file(file_bytes: bytes, filename: str) -> str:
    """
    Extracts high-fidelity structured text (Markdown) from uploaded file bytes.
    Uses Microsoft MarkItDown as the primary engine to preserve tables, headers,
    and lists, with graceful fallbacks to pdfplumber/python-docx/pandas.
    """
    ext = filename.split('.')[-1].lower() if '.' in filename else ''
    
    # 1. 优先使用 Microsoft MarkItDown，完整提取高保真 Markdown（表格、列表、层级标题）
    if _markitdown_instance is not None:
        try:
            stream = io.BytesIO(file_bytes)
            res = _markitdown_instance.convert_stream(stream, file_extension=f".{ext}" if ext else None)
            if res and res.text_content and res.text_content.strip():
                print(f"[file_parser] [SUCCESS] MarkItDown parsed '{filename}' -> {len(res.text_content)} chars Markdown")
                return res.text_content
        except Exception as md_err:
            print(f"[file_parser] [FALLBACK] MarkItDown error on '{filename}': {md_err}, falling back to native parser")

    # 2. 原生/传统解析退坡保障 (Fallbacks)
    if ext == 'pdf':
        return _parse_pdf(file_bytes)
    elif ext in ['doc', 'docx']:
        return _parse_docx(file_bytes)
    elif ext in ['xlsx', 'xls']:
        return _parse_excel(file_bytes)
    elif ext == 'csv':
        try:
            return file_bytes.decode('utf-8')
        except UnicodeDecodeError:
            return file_bytes.decode('gbk', errors='ignore')
    elif ext == 'txt':
        return file_bytes.decode('utf-8', errors='ignore')
    else:
        try:
            return file_bytes.decode('utf-8')
        except:
            raise ValueError(f"Unsupported file format: {ext}")

def _parse_pdf(file_bytes: bytes) -> str:
    text = ""
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as e:
        raise ValueError(f"Failed to parse PDF: {str(e)}")
    return text

def _parse_docx(file_bytes: bytes) -> str:
    try:
        doc = Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
        # 兼容读取表格中的信息
        for table in doc.tables:
            for row in table.rows:
                row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_vals:
                    full_text.append(" | ".join(row_vals))
        return "\n".join(full_text)
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX: {str(e)}")

def _parse_excel(file_bytes: bytes) -> str:
    try:
        import pandas as pd
        excel_data = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
        output = []
        for sheet_name, df in excel_data.items():
            output.append(f"### Sheet: {sheet_name}\n" + df.to_markdown(index=False))
        return "\n\n".join(output)
    except Exception as e:
        raise ValueError(f"Failed to parse Excel: {str(e)}")

def extract_images_from_pdf(file_bytes: bytes) -> list[str]:
    base64_images = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                img = page.to_image(resolution=150)
                img_io = io.BytesIO()
                img.original.save(img_io, format='PNG')
                img_b64 = base64.b64encode(img_io.getvalue()).decode('utf-8')
                base64_images.append(img_b64)
    except Exception as e:
        raise ValueError(f"Failed to extract images from PDF: {str(e)}")
    return base64_images
