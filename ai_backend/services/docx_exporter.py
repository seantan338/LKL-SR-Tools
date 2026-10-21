import os
import copy
import re
import docx
from io import BytesIO
from typing import Dict
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def is_valid_value(val) -> bool:
    """Check if value is present and not a dummy/empty placeholder."""
    if val is None:
        return False
    s = str(val).strip()
    if not s:
        return False
    lower_s = s.lower()
    if lower_s in (
        '[not specified]', 'not specified', 'null', 'none',
        '[未注明]', '未注明', 'n/a', 'na', '-', 'nil', 'unknown'
    ):
        return False
    return True


def add_calibri_run(p, text, bold=False):
    run = p.add_run(text)
    run.font.name = 'Calibri'
    if bold:
        run.bold = True
    if run._element.rPr is not None and run._element.rPr.rFonts is not None:
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    return run


def append_clean_val(p, val, bold=False) -> bool:
    if not is_valid_value(val):
        return False
    val_str = str(val).strip()
    text = p.text
    if text.endswith(' '):
        add_calibri_run(p, val_str, bold=bold)
    else:
        add_calibri_run(p, f" {val_str}", bold=bold)
    return True


def format_qualification(qual_str: str = '', degree_str: str = '', major_str: str = '') -> str:
    """Format and polish qualification string with correct academic prepositions."""
    text = (qual_str or '').strip()
    if not text:
        d = (degree_str or '').strip()
        m = (major_str or '').strip()
        if d and m:
            text = f"{d} {m}".strip()
        else:
            text = d or m or ''

    if not text or not is_valid_value(text):
        return ''

    # 1. Diploma / Certificate / Professional Qualification followed by major without 'in'
    p1 = r'^(Professional\s+Qualification|Higher\s+Diploma|Advanced\s+Diploma|Professional\s+Diploma|Postgraduate\s+Diploma|Diploma|Certificate|Cert)\s+(?!(?:in|of|for)\b)(.+)$'
    m1 = re.match(p1, text, flags=re.IGNORECASE)
    if m1:
        prefix = m1.group(1).title()
        rest = m1.group(2).strip()
        text = f"{prefix} in {rest}"

    # 2. Bachelor / Master / Doctorate followed by major without 'of' or 'in'
    p2 = r'^(Bachelor|Master|Doctorate|Doctor)\s+(?!(?:of|in)\b)(.+)$'
    m2 = re.match(p2, text, flags=re.IGNORECASE)
    if m2:
        prefix = m2.group(1).title()
        rest = m2.group(2).strip()
        text = f"{prefix} of {rest}"

    return text


def create_cloned_bullet_para(prototype_elem, doc, text: str):
    """Clone native template ListParagraph element with numPr and insert text cleanly."""
    new_p_elem = copy.deepcopy(prototype_elem)
    for child in list(new_p_elem):
        if not child.tag.endswith('pPr'):
            new_p_elem.remove(child)
    new_p = Paragraph(new_p_elem, doc)
    add_calibri_run(new_p, text, bold=False)
    return new_p


def export_cv_to_docx(cv_data: Dict, template_path: str) -> BytesIO:
    """
    Takes parsed CV JSON data and injects it into the Sunrise CV docx template.
    Aligns 1:1 with standard Sunrise headhunter deliverables:
    - Bold hierarchies for title, company, position, period
    - Tabular alignment and colon spacing normalization
    - Native Word numPr list paragraphs for responsibilities and skills (automatic hanging indent)
    - Dynamic pruning of empty optional fields (preventing lonely colon lines)
    - Preposition normalization for academic qualifications
    - Multi-blank line collapsing for a tight, elegant layout
    """
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"CV Template not found at {template_path}")

    doc = docx.Document(template_path)

    parsed = cv_data.get('parsed_cv', {})
    if not parsed:
        parsed = cv_data  # fallback if passed directly

    personal = parsed.get('personal_information', {})
    prof_history = parsed.get('professional_history', [])
    education = parsed.get('education', [])
    skills = parsed.get('skills', [])
    languages = parsed.get('languages', [])

    # Extract bullet prototype paragraph from template (Aaaaa)
    bullet_proto = None
    for p in doc.paragraphs:
        if p.text.strip() == 'Aaaaa':
            bullet_proto = copy.deepcopy(p._p)
            for child in list(bullet_proto):
                if not child.tag.endswith('pPr'):
                    bullet_proto.remove(child)
            break

    # Format skills clean list
    if isinstance(skills, str):
        skills_list = [s.strip() for s in skills.split('\n') if s.strip()]
    else:
        skills_list = list(skills)
    clean_skills = [
        s.lstrip('•-* \t').strip() for s in skills_list 
        if is_valid_value(s)
    ]

    # Format languages clean list
    if isinstance(languages, str):
        lang_list = [l.strip() for l in languages.split('\n') if l.strip()]
    else:
        lang_list = [
            str(l).strip() for l in languages 
            if is_valid_value(l)
        ]

    prof_idx = -1
    edu_idx = -1
    current_section = 'HEADER'  # HEADER, PROF, EDU, SKILLS, LANG
    to_delete = []

    skills_para_count = 0
    lang_para_count = 0

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        raw_text = p.text

        # Detect section switches
        if raw_text.startswith('Professional History'):
            current_section = 'PROF'
            continue
        elif raw_text.startswith('Education'):
            current_section = 'EDU'
            continue
        elif text == 'Skills':
            current_section = 'SKILLS'
            continue
        elif text == 'Languages':
            current_section = 'LANG'
            continue

        # --- SECTION: HEADER (Position Applied & Personal Information) ---
        if current_section == 'HEADER':
            if raw_text.startswith('Position Applied for:'):
                pos = parsed.get('position_applied')
                if is_valid_value(pos):
                    append_clean_val(p, pos, bold=True)
            elif raw_text.startswith('Name\t\t\t:'):
                append_clean_val(p, personal.get('name'), bold=False)
            elif raw_text.startswith('Address\t\t\t:'):
                if not append_clean_val(p, personal.get('address'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Age\t\t\t:'):
                if not append_clean_val(p, personal.get('age'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Nationality\t\t:'):
                if not append_clean_val(p, personal.get('nationality'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Race\t\t\t:'):
                if not append_clean_val(p, personal.get('race'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Gender\t\t\t:'):
                if not append_clean_val(p, personal.get('gender'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Marital Status\t\t:'):
                if not append_clean_val(p, personal.get('marital_status'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Last Drawn Salary\t:'):
                if not append_clean_val(p, personal.get('last_drawn_salary'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Expected Salary\t\t:'):
                if not append_clean_val(p, personal.get('expected_salary'), bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Availability\t\t:'):
                if not append_clean_val(p, personal.get('availability'), bold=False):
                    to_delete.append(p)
            continue

        # --- SECTION: PROF_HISTORY ---
        elif current_section == 'PROF':
            if raw_text.startswith('Period\t\t\t:'):
                prof_idx += 1
                if prof_idx >= len(prof_history):
                    to_delete.append(p)
                else:
                    append_clean_val(p, prof_history[prof_idx].get('period'), bold=True)
                continue

            if prof_idx >= len(prof_history):
                to_delete.append(p)
                continue

            if raw_text.startswith('Name of Company\t:'):
                append_clean_val(p, prof_history[prof_idx].get('company'), bold=True)
            elif raw_text.startswith('Position \t\t:') or raw_text.startswith('Position \t:') or raw_text.startswith('Position\t\t:'):
                p.text = ''
                add_calibri_run(p, "Position\t\t: ", bold=True)
                pos = prof_history[prof_idx].get('position')
                if pos:
                    add_calibri_run(p, str(pos).strip(), bold=True)
            elif raw_text.startswith('Last Drawn Salary\t:'):
                sal = prof_history[prof_idx].get('last_drawn_salary')
                if not append_clean_val(p, sal, bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Reason for Leaving\t:'):
                rfl = prof_history[prof_idx].get('reason_for_leaving')
                if not append_clean_val(p, rfl, bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Job Responsibilities\t:'):
                p.text = ''
                add_calibri_run(p, "Job Responsibilities\t:", bold=False)
            elif text == 'Aaaaa':
                p.text = ''
                resps = prof_history[prof_idx].get('responsibilities', [])
                if isinstance(resps, str):
                    resps = [r.strip() for r in resps.split('\n') if r.strip()]
                clean_resps = [
                    r.lstrip('•-* \t').strip() for r in resps 
                    if is_valid_value(r)
                ]
                if clean_resps:
                    # First item uses current template paragraph (preserves numPr)
                    add_calibri_run(p, clean_resps[0], bold=False)
                    prev_p = p
                    # Subsequent items clone native bullet prototype paragraph
                    for extra_r in clean_resps[1:]:
                        if bullet_proto is not None:
                            new_p = create_cloned_bullet_para(bullet_proto, doc, extra_r)
                        else:
                            new_p = doc.add_paragraph()
                            add_calibri_run(new_p, extra_r, bold=False)
                        prev_p._element.addnext(new_p._element)
                        prev_p = new_p
                else:
                    to_delete.append(p)
            elif text in ['Bbbbb', 'Cccccc']:
                to_delete.append(p)
            elif raw_text == ' ':
                p.text = ''
            continue

        # --- SECTION: EDUCATION ---
        elif current_section == 'EDU':
            if raw_text.startswith('Year\t\t\t:'):
                edu_idx += 1
                if edu_idx >= len(education):
                    to_delete.append(p)
                else:
                    yr = education[edu_idx].get('year')
                    if not append_clean_val(p, yr, bold=False):
                        to_delete.append(p)
                continue

            if edu_idx >= len(education):
                to_delete.append(p)
                continue

            if raw_text.startswith('Name of School\t\t:'):
                sch = education[edu_idx].get('school')
                if not append_clean_val(p, sch, bold=False):
                    to_delete.append(p)
            elif raw_text.startswith('Qualification\t:') or raw_text.startswith('Qualification\t\t:'):
                qual_raw = education[edu_idx].get('qualification')
                deg_raw = education[edu_idx].get('degree')
                maj_raw = education[edu_idx].get('major')
                formatted_qual = format_qualification(qual_raw, deg_raw, maj_raw)
                p.text = ''
                add_calibri_run(p, "Qualification\t\t: ", bold=False)
                if formatted_qual:
                    add_calibri_run(p, formatted_qual, bold=False)
            continue

        # --- SECTION: SKILLS ---
        elif current_section == 'SKILLS':
            skills_para_count += 1
            if skills_para_count == 1:
                p.text = ''
                if clean_skills:
                    if bullet_proto is not None:
                        prev_p = p
                        for s_item in clean_skills:
                            new_p = create_cloned_bullet_para(bullet_proto, doc, s_item)
                            prev_p._element.addnext(new_p._element)
                            prev_p = new_p
                        to_delete.append(p)
                    else:
                        add_calibri_run(p, '\n'.join([f"• {s}" for s in clean_skills]), bold=False)
                else:
                    to_delete.append(p)
            else:
                to_delete.append(p)
            continue

        # --- SECTION: LANGUAGES ---
        elif current_section == 'LANG':
            lang_para_count += 1
            if lang_para_count == 1:
                p.text = ''
                if lang_list:
                    add_calibri_run(p, lang_list[0], bold=False)
                    prev_p = p
                    for extra_lang in lang_list[1:]:
                        new_p = doc.add_paragraph()
                        add_calibri_run(new_p, extra_lang, bold=False)
                        prev_p._element.addnext(new_p._element)
                        prev_p = new_p
                else:
                    to_delete.append(p)
            else:
                to_delete.append(p)
            continue

    # Delete marked paragraphs
    for p in to_delete:
        elem = p._element
        if elem.getparent() is not None:
            elem.getparent().remove(elem)

    # Post-process 1: Merge consecutive blank lines across the entire document
    paras = list(doc.paragraphs)
    prev_empty = False
    for p in paras:
        is_empty = (p.text.strip() == '')
        if is_empty:
            if prev_empty:
                elem = p._element
                if elem.getparent() is not None:
                    elem.getparent().remove(elem)
            else:
                prev_empty = True
        else:
            prev_empty = False

    # Post-process 2: Ensure exactly 1 blank line before major section headings (e.g. Languages, Skills, Education)
    section_titles = {'Personal Information', 'Professional History', 'Education', 'Education ', 'Skills', 'Languages'}
    # First, trim extra blank lines before headings
    paras = list(doc.paragraphs)
    for i in range(len(paras)):
        if paras[i].text.strip() in section_titles and i > 0:
            j = i - 1
            blank_count = 0
            while j >= 0 and paras[j].text.strip() == '':
                blank_count += 1
                if blank_count > 1:
                    elem = paras[j]._element
                    if elem.getparent() is not None:
                        elem.getparent().remove(elem)
                j -= 1

    # Second, ensure at least 1 blank line before headings if not at top of document
    paras = list(doc.paragraphs)
    for i in range(1, len(paras)):
        if paras[i].text.strip() in section_titles:
            if paras[i - 1].text.strip() != '':
                paras[i].insert_paragraph_before('')

    # Post-process 3: Strip trailing blank paragraphs at document end
    paras = list(doc.paragraphs)
    while paras and paras[-1].text.strip() == '':
        elem = paras[-1]._element
        if elem.getparent() is not None:
            elem.getparent().remove(elem)
        paras.pop()

    out_stream = BytesIO()
    doc.save(out_stream)
    out_stream.seek(0)
    return out_stream

